"""Slide deck in the team's template format (56940c12 MIMIC-IIINote.pptx): keeps the template's TabPFN / UTDE slides
(section 1) and adds section 2 (late fusion, TabPFN-in-UTDE, fair protocol) and section 3 (results, SHAP, appendix)
in the same card / header-bar styles.  Usage: python scripts/28_make_slides_template.py <template.pptx>
"""
import re, sys, copy
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT = Path(__file__).resolve().parents[1]; R = ROOT / 'results'
TPL = Path(sys.argv[1]); OUT = R / 'slides' / 'MIMIC_TabPFN_UTDE_fusion_template_format.pptx'
ARROW = Path('/tmp/claude-1010/-mnt-nas5-AIBL-Research-hskim/553fe755-660f-4865-bfa0-4f6b37b5b62a/scratchpad/tpl/arrow_5_Image_0.png')
F = 'IBM Plex Sans KR'
C = dict(ink='131B22', body='3B4652', blue='1D5FB4', blue2='2E6FC4', paleblue='B9CBE4', lightblue='DCE6F3', card='E6EBEA', paper='F3F5F4',
         amber='8A5E06', amberfill='E2A72E', paleamber='F2E3C0', ref='5E6873', dim='C9D1D8', navy='1F3A5F', navytxt='C8D3E6', red='C0392B')
rgb = lambda k: RGBColor.from_string(C.get(k, k))

prs = Presentation(TPL)
sld = prs.slides._sldIdLst; N_TPL = len(list(sld))        # template slides are dropped *after* the new ones are added (part-name collisions otherwise)
L_TB = [l for l in prs.slide_layouts if l.name == 'Title and body'][0]; L_BLANK = [l for l in prs.slide_layouts if l.name == 'Blank'][0]


def rich(tf, text, size=12, color='ink', bold_color=None, align=PP_ALIGN.LEFT, first=True):
    """'**bold**' segments become bold (optionally coloured); '\n' starts new paragraphs."""
    for k, line in enumerate(text.split('\n')):
        p = tf.paragraphs[0] if (first and k == 0) else tf.add_paragraph(); p.alignment = align
        for j, seg in enumerate(re.split(r'(\*\*.*?\*\*)', line)):
            if not seg:
                continue
            b = seg.startswith('**'); r = p.add_run(); r.text = seg.strip('*') if b else seg
            r.font.name = F; r.font.size = Pt(size); r.font.bold = b; r.font.color.rgb = rgb(bold_color if (b and bold_color) else color)


def tbox(s, x, y, w, h, text, size=12, color='ink', bold_color=None, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    t = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = t.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.04); tf.margin_top = tf.margin_bottom = Inches(0.02); rich(tf, text, size, color, bold_color, align); return t


def card(s, x, y, w, h, fill, text='', size=12, color='ink', bold_color=None, corner=0.03, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); sh.adjustments[0] = corner
    sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill); sh.line.fill.background(); sh.shadow.inherit = False
    tf = sh.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor; tf.margin_left = tf.margin_right = Inches(0.12); tf.margin_top = tf.margin_bottom = Inches(0.08)
    if text:
        rich(tf, text, size, color, bold_color, align)
    return sh


def arrow(s, x, y):
    s.shapes.add_picture(str(ARROW), Inches(x), Inches(y), width=Inches(0.25), height=Inches(0.12))


def tb_slide(title, subtitle, sub_color='blue'):
    s = prs.slides.add_slide(L_TB)
    for ph in list(s.placeholders):
        if ph.placeholder_format.idx == 0:
            ph.text_frame.text = ''; r = ph.text_frame.paragraphs[0].add_run(); r.text = title; r.font.name = F; r.font.size = Pt(30); r.font.bold = True; r.font.color.rgb = rgb('ink')
        else:
            ph._element.getparent().remove(ph._element)
    tbox(s, 0.89, 1.16, 11.9, 0.3, subtitle, 13, 'body', sub_color)
    return s


def bar_slide(section, title, lead=None):
    s = prs.slides.add_slide(L_BLANK)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.9)); bar.fill.solid(); bar.fill.fore_color.rgb = rgb('navy'); bar.line.fill.background(); bar.shadow.inherit = False
    t = s.shapes.add_textbox(Inches(0.4), Inches(0.08), Inches(9.5), Inches(0.3)); t.text_frame.word_wrap = True; pp = t.text_frame.paragraphs[0]; pp.alignment = PP_ALIGN.LEFT; r = pp.add_run(); r.text = section; r.font.name = 'Malgun Gothic'; r.font.size = Pt(12); r.font.color.rgb = rgb('navytxt')
    t = s.shapes.add_textbox(Inches(0.4), Inches(0.36), Inches(12.4), Inches(0.47)); t.text_frame.word_wrap = True; pp = t.text_frame.paragraphs[0]; pp.alignment = PP_ALIGN.LEFT; r = pp.add_run(); r.text = title; r.font.name = 'Malgun Gothic'; r.font.size = Pt(22); r.font.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    if lead:
        tbox(s, 0.6, 1.0, 12.1, 0.5, lead, 13, 'body', 'blue')
    return s


def table(s, x, y, w, rows, col_w, size=11, bold_rows=(), red=()):
    shape = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(0.3 * len(rows))); t = shape.table
    tblPr = shape._element.graphic.graphicData.tbl.tblPr; tblPr.set('bandRow', '0'); tblPr.set('firstRow', '0')
    for j, cw in enumerate(col_w):
        t.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j); c.text = ''; p = c.text_frame.paragraphs[0]; r = p.add_run(); r.text = str(val); r.font.name = F; r.font.size = Pt(size)
            r.font.bold = i == 0 or i in bold_rows; r.font.color.rgb = rgb('red') if (i, j) in red else rgb('ink')
            c.fill.solid(); c.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if i else rgb('card'); c.margin_left = c.margin_right = Inches(0.06); c.margin_top = c.margin_bottom = Inches(0.02)
    return t


def ref(s, text):
    tbox(s, 0.89, 6.82, 11.9, 0.3, text, 12, 'ref')


# ================================================================ 0. setup (card style)
s = tb_slide('Task · Cohort · Protocol', '**28일 사망 예측** · ICU 입실 후 첫 W시간(6 h / 24 h)의 자료만 사용 → 창 안에서 사망한 stay는 제외(landmark) → PATIENTS.DOD로 28일 사망 판정')
card(s, 0.89, 1.8, 2.5, 2.3, 'card', '**① Cohort**\n성인 → SOFA 6개 장기 측정 → lactate·BE 있음 → W시점 생존\n\n**6 h 5,757** stays (사망 29 %)\n**24 h 9,229** stays (26 %)', 12, 'ink', 'blue')
arrow(s, 3.5, 2.85)
card(s, 3.86, 1.8, 2.9, 2.3, 'card', '**② Split**\n환자 단위 70 / 15 / 15\n모든 모델이 같은 분할·레이블 사용\n\ntest **853 / 1,374** stays', 12, 'ink', 'blue')
arrow(s, 6.87, 2.85)
card(s, 7.23, 1.8, 2.9, 2.3, 'blue', '**③ Inputs (통일 변수)**\nSOFA 입력(PaO2, FiO2, platelets, bilirubin, creatinine, MAP, GCS, 승압제, 소변량, 환기)\n+ lactate · base excess\n+ age · sex', 12, 'paper', 'paper')
arrow(s, 10.24, 2.85)
card(s, 10.6, 1.8, 1.85, 2.3, 'ink', '**④ Metric**\ntest AUROC (AUPRC)\n\npaired bootstrap\nΔAUROC [95 % CI]\nvs TabPFN', 12, 'dim', 'paper')
tbox(s, 0.89, 4.49, 5.1, 0.3, 'Leakage controls', 15, 'blue', 'blue')
tbox(s, 0.89, 4.83, 5.1, 1.7, '· 노트: INTIME + W 이전에 **저장(STORETIME)** 된 것만\n· 검사: 입실 −6 h ~ +W h, 차트·투약·배출량: 0 ~ W h\n· 스태커·임계값은 validation에서만 적합', 12)
tbox(s, 6.24, 4.49, 6.4, 0.3, 'Selection check', 15, 'amber', 'amber')
tbox(s, 6.24, 4.83, 6.4, 1.7, '· **SOFA 계산 가능** 여부가 지배적 선택: 사망률 24.0 % vs 10.0 % (보정 OR 1.9)\n· 그 안에서 lactate 유무는 OR 1.2\n· 선택군에서 예측이 더 쉬워지지 않음: 같은 공변량 모델 AUROC 0.70 vs 0.76', 12)
ref(s, 'MIMIC-III v1.4 (Johnson et al., 2016). 결과: results/selection_bias/, results/flowchart/')

# ================================================================ 2a. late fusion (card style)
s = tb_slide('Late Fusion (TabPFN + UTDE)', '**Multimodal (후기 결합)** · 두 모델을 따로 학습한 뒤 **최종 로짓 2개**만 로지스틱 스태커로 결합 → 임베딩은 섞지 않음')
card(s, 0.89, 1.8, 2.6, 2.3, 'lightblue', '**① TabPFN**\n구조화 68열\n→ p_T\n\n검증 확률: train만 context\n테스트 확률: train + val context', 12, 'ink', 'blue')
card(s, 0.89, 4.2, 2.6, 1.6, 'paleamber', '**① UTDE**\n시계열 + 노트 → p_U\n\n검증 로짓: 선택된 체크포인트에서', 12, 'ink', 'amber')
arrow(s, 3.6, 2.85); arrow(s, 3.6, 4.9)
card(s, 3.96, 1.8, 3.9, 4.0, 'blue', '**② Stacker (validation에서만 적합)**', 13, 'paper', 'paper')
card(s, 4.1, 2.35, 3.6, 1.0, 'blue2', '**z = logit(p)** · 확률 → 로짓\n극단값이 눌리지 않고, 두 모델의 보정 차이를 계수가 흡수', 12, 'paper', 'paper')
card(s, 4.1, 3.45, 3.6, 1.05, 'blue2', '**p_fusion = σ(β₀ + β₁ z_T + β₂ z_U)**\n파라미터 3개, sklearn LogisticRegression (C = 1)', 12, 'paper', 'paper')
card(s, 4.1, 4.6, 3.6, 1.05, 'blue2', '**학습된 계수** · 6 h 0.55 / 0.54, 24 h 0.41 / 0.43\n→ 거의 동일 가중; 확률 단순 평균과 결과 ≈ 같음', 12, 'paper', 'paper')
arrow(s, 7.96, 3.7)
card(s, 8.32, 1.8, 4.1, 4.0, 'ink', '**③ Output · test에 계수 그대로 적용**', 13, 'paper', 'paper')
tbox(s, 8.5, 2.4, 3.8, 0.5, '**P(28일 사망)**', 15, 'paper', 'paper')
tbox(s, 8.5, 2.95, 3.8, 2.7, '21채널 UTDE: **0.845 / 0.860**\n(+0.014 ns / +0.023 [+0.010, +0.035])\n\n통일 변수 UTDE-SOFA: **0.839 / 0.841**\n(+0.008 / +0.003, ns)\n\n→ 이득의 대부분은 UTDE만 보던 활력징후 변수', 12, 'dim')
tbox(s, 0.89, 6.1, 11.9, 0.6, '**WHY?** ① 테스트 레이블 미사용 · ② 두 모델의 확신 정도(로짓 분산)를 계수가 맞춰 줌 · ③ 가장 단순한 멀티모달 기준선 — 복잡한 결합이 이것을 넘는지가 검증 대상', 12, 'ink', 'amber')
ref(s, 'scripts/14_late_fusion.py · paired bootstrap 2,000회')

# ================================================================ 2b. TabPFN-in-UTDE (card style)
s = tb_slide('TabPFN-in-UTDE (replace / 3-stream)', '**Multimodal (단일 네트워크)** · frozen TabPFN의 내부 표현을 **토큰 8개**로 꺼내 UTDE의 attention에 넣고 end-to-end 학습')
card(s, 0.89, 1.8, 2.6, 2.6, 'card', '**① TabPFN embedding**\n앙상블 멤버 8 × 192-d (query 모드)\ntrain 행은 5-fold context로 자기 자신 제외\n\n선형 프로브 0.834 ≈ TabPFN 0.837 → 정보 보존', 12, 'ink', 'blue')
arrow(s, 3.6, 3.0)
card(s, 3.96, 1.8, 4.1, 2.6, 'blue', '**② replace**', 13, 'paper', 'paper')
card(s, 4.1, 2.3, 3.8, 0.55, 'blue2', '**TabPFN 토큰 8** (Linear 192→128)', 12, 'paper', 'paper')
card(s, 4.1, 2.95, 3.8, 0.55, 'blue2', '노트 토큰 5 (Clinical-Longformer)', 12, 'paper', 'paper')
card(s, 4.1, 3.6, 3.8, 0.55, 'blue2', 'cross-attention → MLP  (시계열 stream 없음)', 12, 'paper', 'paper')
arrow(s, 8.16, 3.0)
card(s, 8.52, 1.8, 3.9, 2.6, 'ink', '**② 3-stream**', 13, 'paper', 'paper')
card(s, 8.66, 2.3, 3.6, 0.5, 'paleblue', '시계열 토큰 tt_max (mTAND + grid)', 11, 'ink')
card(s, 8.66, 2.85, 3.6, 0.5, 'paleblue', '**+ TabPFN 토큰 8** (이어 붙임, 위치표 tt_max + 8)', 11, 'ink', 'blue')
card(s, 8.66, 3.4, 3.6, 0.5, 'paleblue', '노트 토큰 5 → cross-attention → MLP', 11, 'ink')
tbox(s, 8.66, 3.95, 3.6, 0.4, '학습 일정·손실·lr은 UTDE 원본과 동일', 11, 'dim')
tbox(s, 0.89, 4.65, 5.1, 0.3, 'Implementation', 15, 'blue', 'blue')
tbox(s, 0.89, 4.99, 5.1, 1.6, '· UTDE 포크 플래그 --tab_emb, --tab_emb_mode {replace, concat}\n· 6 epoch, Longformer 3회 갱신, lr 4e-4 / 2e-5, 검증 AUROC로 체크포인트\n· 변수 통일: 시계열 15채널(SOFA 입력 + lactate·BE + 환기 + age/sex)', 12)
tbox(s, 6.24, 4.65, 6.4, 0.3, 'WHY?', 15, 'amber', 'amber')
tbox(s, 6.24, 4.99, 6.4, 1.6, '① **단일 모델**로 늦은 융합과 같거나 나은지 — 배포·서사가 단순해짐\n② 노트가 TabPFN 표현을 **조건으로** 읽히면 선형 스태커가 못 잡는 상호작용 가능\n③ replace vs 3-stream 비교로 "시계열 stream이 필요한가"에 답', 12)
ref(s, 'scripts/23_tabpfn_embeddings.py · MultimodalMIMIC 패치 (run_tabemb.sh)')

# ================================================================ 2c. fair protocol (bar + table)
s = bar_slide('2. Proposed methods', 'Fair comparison protocol — 변수를 맞춘 뒤 구조를 비교', '원 UTDE 입력(21채널)에는 **HR·RR·SpO2·체온·혈압·O2 flow**가 있었고 TabPFN(SOFA 입력)에는 없었음 → 모델 비교에 변수 차이가 섞임')
table(s, 0.6, 1.6, 7.2, [['', 'TabPFN (68)', 'UTDE 21 ch', 'UTDE-SOFA 15 ch', 'TabPFN + vitals'],
    ['SOFA 입력 + lactate / BE', '집계', '시계열', '시계열', '집계'], ['age / sex', '○', '×', '상수 채널', '○'], ['환기 (SOFA 호흡 점수 입력)', 'vent_n', '×', '0/1 채널', '○'],
    ['HR·RR·SpO2·Temp·혈압·O2flow', '×', '○ (9채널)', '×', '○ (집계 45열)'], ['노트', '×', '○', '○', '×']], [2.6, 1.1, 1.1, 1.3, 1.1], 11)
tbox(s, 8.1, 1.6, 4.6, 0.3, '통제 1 — UTDE-SOFA', 15, 'blue', 'blue')
tbox(s, 8.1, 1.95, 4.6, 1.2, 'UTDE 채널을 TabPFN의 변수로 제한: SOFA 입력 10 + 환기 + lactate + BE + age + sex = **15채널**; 노트 stream은 그대로', 12)
tbox(s, 8.1, 3.2, 4.6, 0.3, '통제 2 — TabPFN + vitals', 15, 'amber', 'amber')
tbox(s, 8.1, 3.55, 4.6, 1.2, '반대로 TabPFN에 활력징후 9종의 min·max·mean·last·count **45열**을 추가 (113열)', 12)
tbox(s, 0.6, 4.2, 7.2, 2.4, '**평가** · 같은 테스트셋, paired bootstrap 2,000회, Δ vs TabPFN(68) 또는 vs TabPFN + vitals\n**해석** · 세 모델 모두에 어트리뷰션 — TabPFN permutation SHAP, UTDE occlusion(모달리티·채널·TabPFN 토큰), 스태커 정확 SHAP + feature-level 분해\n**기대** · 변수를 맞추면 "구조의 효과"만 남음: 멀티모달 이득이 남으면 노트·궤적의 가치, 사라지면 변수의 가치였던 것', 12)
tbox(s, 0.6, 6.85, 12, 0.3, 'scripts/13_build_utde_dataset.py --channel-set sofa · 25_tabpfn_vitals.py', 10, 'ref')

# ================================================================ 3a0. unimodal tier table with CIs (template slide-8 style)
import json as _json
U = _json.load(open(R / 'sofa_tiers_demo' / 'unimodal_table_ci.json'))
f3 = lambda d, k: f"{d[k]:.3f} [{d[k+'_ci95'][0]:.3f}, {d[k+'_ci95'][1]:.3f}]"
s = bar_slide('3. Result and Discussion', 'Feature selection (Unimodal) — test AUROC [95 % bootstrap CI]')
tbox(s, 0.6, 1.05, 11.9, 0.3, 'AUROC [95 % CI] · 6 h n = 853, 24 h n = 1,374 · 1,000 bootstrap resamples', 13, 'amber', 'amber')
rows = [['tier', 'XGB 6h', 'TabPFN 6h', 'XGB 24h', 'TabPFN 24h'],
        ['SOFA score 단독', f3(U['sofa_alone_W6'], 'auroc'), '(same)', f3(U['sofa_alone_W24'], 'auroc'), '(same)']]
for demo in (False, True):
    if demo:
        rows.append(['+age/sex', '', '', '', ''])
    for r in [r for r in U['rows'] if r['demo'] == demo]:
        rows.append([r['tier']] + [f"{r[k]['test_auroc']:.3f} [{r[k]['test_auroc_ci95'][0]:.3f}, {r[k]['test_auroc_ci95'][1]:.3f}]" for k in ('xgb_6', 'tabpfn_6', 'xgb_24', 'tabpfn_24')])
table(s, 0.6, 1.4, 12.1, rows, [2.5, 2.4, 2.4, 2.4, 2.4], 11, bold_rows=(6, 10), red={(10, 2), (10, 3), (10, 4)})
tbox(s, 0.6, 5.05, 12.1, 1.7, '· AUPRC [95 % CI] (tier 4 + age/sex): XGB 0.662 [0.604, 0.718] / 0.660 [0.610, 0.706], TabPFN 0.681 [0.621, 0.739] / 0.673 [0.621, 0.717]; SOFA 단독 0.443 [0.386, 0.500] / 0.428 [0.378, 0.481]\n· 신뢰구간 폭 ≈ ±0.03 (6 h) / ±0.02 (24 h): tier 간 차이(≤ 0.02)와 XGB–TabPFN 차이(≤ 0.015)는 모두 CI 안 — **점추정 순위는 참고용**\n· 확실한 것: 학습 모델 vs SOFA 단독 (+0.17~0.18, CI 비중첩), age/sex 추가 (+0.01~0.02, 6 h에서 더 큼)\n· 전체 표(AUPRC 포함): results/sofa_tiers_demo/unimodal_table_ci.md', 11)

# ================================================================ 3a. results tables
s = bar_slide('3. Result and Discussion', 'Quantitative results — unified variables and the fair baseline')
tbox(s, 0.6, 1.05, 11.9, 0.3, 'Unified variable set (SOFA + lactate/BE + age/sex)', 13, 'amber', 'amber')
table(s, 0.6, 1.4, 12.1, [['Method', 'Modality', '6 h AUROC (AUPRC)', 'Δ vs TabPFN', '24 h AUROC (AUPRC)', 'Δ vs TabPFN'],
    ['TabPFN v2 (68 features)', 'Unimodal', '0.832 (.681)', '–', '0.837 (.673)', '–'],
    ['UTDE-SOFA (15 ch + notes)', 'Multimodal', '0.823 (.658)', '−0.009 [−0.031, +0.013]', '0.826 (.658)', '−0.011 [−0.025, +0.002]'],
    ['Late fusion (TabPFN + UTDE-SOFA)', 'Multimodal', '0.839 (.688)', '+0.008 [−0.008, +0.023]', '0.841 (.680)', '+0.003 [−0.002, +0.009]'],
    ['TabPFN-in-UTDE, replace', 'Multimodal', '0.827 (.677)', '−0.005 [−0.012, +0.002]', '0.836 (.671)', '−0.001 [−0.006, +0.004]'],
    ['TabPFN-in-UTDE, 3-stream', 'Multimodal', '0.835 (.687)', '+0.004 [−0.010, +0.016]', '0.856 (.709)', '+0.019 [+0.008, +0.030]']],
    [3.0, 1.1, 1.6, 2.5, 1.6, 2.5], 10, bold_rows=(5,), red={(3, 2), (5, 4), (5, 5)})
tbox(s, 0.6, 3.75, 11.9, 0.3, 'With the vital-sign channels — fair structured baseline = TabPFN + vitals', 13, 'amber', 'amber')
table(s, 0.6, 4.1, 12.1, [['Method', '6 h', 'Δ vs TabPFN+vitals', '24 h', 'Δ vs TabPFN+vitals'],
    ['TabPFN + vitals aggregates (113 features)', '0.842 (.705)', '–', '0.852 (.684)', '–'],
    ['UTDE 21 ch + notes', '0.798 (.618)', '−0.044 [−0.072, −0.016]', '0.844 (.680)', '−0.008 [−0.026, +0.010]'],
    ['Late fusion (TabPFN+vitals + UTDE 21 ch)', '0.844 (.710)', '+0.002 [−0.012, +0.016]', '0.863 (.709)', '+0.012 [+0.002, +0.021]'],
    ['TabPFN-in-UTDE, 3-stream (21 ch)', '0.836 (.700)', '−0.006 [−0.017, +0.004]', '0.868 (.731)', '+0.016 [+0.003, +0.028]']],
    [3.8, 1.4, 2.75, 1.4, 2.75], 10, bold_rows=(1, 4), red={(4, 3), (4, 4)})
tbox(s, 0.6, 5.85, 12.1, 1.1, '· 변수를 맞추면 늦은 융합·replace는 TabPFN을 넘지 못함; **3-stream만 24 h에서 유의** (+0.019; 늦은 융합 대비 +0.016 [+0.006, +0.026])\n· 활력징후 집계만 TabPFN에 더해도 +0.010 / +0.015 — 공정 기준선 위의 멀티모달 순이득은 **+0.01~0.02, 24 h에서만** · test n = 853 / 1,374 · single seed', 11)

# ================================================================ 3b. SHAP headline (template slide-9 style: two images)
s = bar_slide('3. Result and Discussion', 'Shap value', '24 h · 변수를 맞추면 모든 모델이 **GCS·age·lactate·소변량**에 의존 — 구조가 바꾸는 것은 "무엇을 보느냐"가 아니라 "어떻게 합치느냐"')
s.shapes.add_picture(str(R / 'fusion_shap/tabpfn_shap_W24.png'), Inches(0.2), Inches(1.6), width=Inches(6.9))
s.shapes.add_picture(str(R / 'unified_shap/fusion_sofa_feature_beeswarm_W24.png'), Inches(7.6), Inches(1.6), height=Inches(5.3))
tbox(s, 0.2, 5.3, 6.9, 0.4, 'TabPFN (68 features) — permutation SHAP, 24 h', 12, 'ink', align=PP_ALIGN.CENTER)
tbox(s, 7.6, 6.95, 5.5, 0.35, 'Late fusion TabPFN + UTDE-SOFA — feature-level, 24 h', 12, 'ink', align=PP_ALIGN.CENTER)
tbox(s, 0.2, 5.75, 6.9, 1.2, '· TabPFN: gcs_last ≫ age > urine_n · gcs_count · lactate · BE (6 h: age ≈ gcs_last)\n· 융합: gcs_last (T) 0.80 ≫ age (T) 0.32 > notes (U) 0.24 > GCS (U) 0.23 — GCS·age가 양 경로로 중복, UTDE 경로의 고유 기여는 노트뿐\n· 측정 횟수 feature(urine_n, gcs_count)가 상위 — 기록 패턴 학습 (limitation)', 11)

s = bar_slide('3. Result and Discussion', 'Shap value — 모델별 어트리뷰션 요약', 'UTDE-SOFA 채널 occlusion(24 h)과 다섯 모델의 요약표')
s.shapes.add_picture(str(R / 'unified_shap/utde_sofa_beeswarm_W24.png'), Inches(0.2), Inches(1.55), height=Inches(4.0))
tbox(s, 0.2, 5.55, 5.7, 0.35, 'UTDE-SOFA — channel occlusion, 24 h (0/1 채널은 창 내 유무로 채색)', 11, 'ink', align=PP_ALIGN.CENTER)
table(s, 6.1, 1.6, 6.9, [['model', 'drives it (24 h)', 'notes removed\n(6 h / 24 h)'],
    ['TabPFN', 'gcs_last ≫ age > urine_n, gcs_count, lactate, BE', '–'],
    ['UTDE-SOFA', 'GCS (−0.057) > Age (−0.020) > lactate, PaO2, bilirubin', '−0.018 / −0.023'],
    ['Late fusion', 'TabPFN 로짓 61 % (6 h 39 %); 로짓 상관 0.91', 'via UTDE logit'],
    ['TabPFN-in-UTDE replace', 'TabPFN 토큰 −0.27; 노트 branch 미사용', '−0.001 / −0.002'],
    ['TabPFN-in-UTDE 3-stream', 'TabPFN 토큰 −0.091 > 노트 −0.016 > 시계열 −0.005', '−0.007 / −0.016']], [1.9, 3.6, 1.4], 10, bold_rows=(5,), red={(5, 2)})
tbox(s, 6.1, 4.3, 6.9, 2.4, '· UTDE-SOFA에서도 GCS·Age가 1·2위 — 같은 변수를 주면 TabPFN과 같은 순서로 의존\n· 6 h에서는 Age ≈ GCS가 공동 1위 (TabPFN 6 h와 동일)\n· 노트 제거 시 손실은 −0.02 수준으로 21채널 UTDE와 같음 — 노트의 몫은 변수 통일과 무관\n· replace vs 3-stream: 노트 사용 여부가 성능 차이(0.836 vs 0.856)를 설명', 11)

# ================================================================ 3c. TabPFN-in-UTDE attribution
s = bar_slide('3. Result and Discussion', 'Shap value — TabPFN-in-UTDE: replace는 노트를 버리고, 3-stream은 노트를 쓴다', 'TabPFN 토큰·노트·시계열 stream을 각각 제거했을 때의 test AUROC (24 h)')
table(s, 0.6, 1.6, 6.0, [['제거 대상 (24 h)', 'replace', '3-stream'], ['없음 (전체)', '0.836', '0.856'], ['TabPFN 토큰', '0.562 (−0.274)', '0.765 (−0.091)'],
    ['노트', '0.834 (−0.002)', '0.840 (−0.016)'], ['시계열 stream', '(없음)', '0.852 (−0.005)'], ['토큰 + 노트', '0.494', '0.752']], [2.4, 1.8, 1.8], 12, red={(3, 2)})
tbox(s, 0.6, 3.6, 6.0, 3.2, '· replace: TabPFN 토큰이 처음부터 완성된 저잡음 신호 → 네트워크가 그것만 통과시키고 노트 branch는 학습 신호를 못 받음(검증 AUROC epoch 2 이후 정체)\n· 3-stream: 시계열 branch가 "아직 배워야 하는" 쪽이라 gradient가 계속 흐르고 노트 branch도 함께 학습 → 노트 −0.016 (6 h −0.007)\n· 시계열 stream의 직접 기여는 −0.005, 채널별 ≤ 0.006 — 정보는 TabPFN 토큰과 중복이지만 **노트 branch를 깨우는 역할**\n· 결론: 같은 정보에서 3-stream의 +0.019는 "TabPFN 표현을 조건으로 노트를 읽는 것"에서 나옴', 12)
s.shapes.add_picture(str(R / 'unified_shap/tabemb3_sofa_beeswarm_W24.png'), Inches(6.9), Inches(1.5), height=Inches(5.2))
tbox(s, 6.9, 6.75, 6.0, 0.3, 'TabPFN-in-UTDE 3-stream (24 h) — channel occlusion (0/1 채널은 창 내 유무로 채색)', 11, 'ink', align=PP_ALIGN.CENTER)

# ================================================================ summary
s = bar_slide('3. Result and Discussion', 'Summary · Limitations · Next steps')
tbox(s, 0.6, 1.1, 6.0, 0.3, 'Summary', 15, 'blue', 'blue')
tbox(s, 0.6, 1.45, 6.0, 4.5, '· 구조화 feature + TabPFN이 가장 강한 단일 표현 (0.832 / 0.837); 같은 변수에서 UTDE·늦은 융합은 이를 넘지 못함\n· 기존 융합 이득(+0.023)의 대부분은 UTDE만 보던 활력징후 변수 — TabPFN에 넣으면 0.852\n· TabPFN을 UTDE 안에 넣을 때 **시계열 stream을 남긴 3-stream**만 유의한 이득 (24 h +0.019; 공정 기준선 대비 +0.016)\n· 어트리뷰션: 모든 모델이 GCS·age·lactate·소변량에 의존; 노트는 24 h에서만, 3-stream에서만 실제로 쓰임', 12)
tbox(s, 6.9, 1.1, 5.8, 0.3, 'Limitations', 15, 'amber', 'amber')
tbox(s, 6.9, 1.45, 5.8, 1.6, '· 단일 seed · 단일 데이터셋(MIMIC-III)\n· 측정 횟수 feature = 진료 패턴 학습\n· 적용 대상: ABG·lactate가 시행된 ICU 환자(SOFA 계산 가능군, 사망률 2.4배)', 12)
tbox(s, 6.9, 3.2, 5.8, 0.3, 'Next steps', 15, 'blue', 'blue')
tbox(s, 6.9, 3.55, 5.8, 2.4, '· seed 3회 반복으로 3-stream vs 늦은 융합(+0.016) 신뢰구간 확정\n· 3-stream에 TabPFN + vitals 토큰 사용 (공정 기준선과 같은 입력)\n· UTDE 노트 수 5 → 10 (24 h stay의 63 %가 노트 5개 초과)\n· Agent 단계: 위험 모델을 도구로, 노트는 질문 기반 검색, 능동 검사 획득 과제', 12)
tbox(s, 0.6, 6.3, 12.1, 0.4, 'github.com/HyeongSubKim/mimic3-sofa28-mortality-multimodal (Apache-2.0) — 재현 가이드 · UTDE 패치 · 결과표 · 그림', 11, 'ref')

# ================================================================ appendix (template slides 10–14 style)
s = prs.slides.add_slide(L_BLANK)
for ph in list(s.placeholders):
    ph._element.getparent().remove(ph._element)
tbox(s, 5.74, 3.48, 1.9, 0.6, '**Appendix**', 20, 'ink', align=PP_ALIGN.CENTER)
for W in (6, 24):
    s = bar_slide('3. Result', f'Shap value- {W}h / Unimodal')
    s.shapes.add_picture(str(R / f'fusion_shap/tabpfn_shap_W{W}.png'), Inches(0.0), Inches(1.35), height=Inches(5.6))
    s.shapes.add_picture(str(R / f'unified_shap/utde_sofa_beeswarm_W{W}.png'), Inches(6.2), Inches(1.35), height=Inches(5.6))
    s = bar_slide('3. Result', f'Shap value- {W}h / Multimodal')
    s.shapes.add_picture(str(R / f'unified_shap/fusion_sofa_feature_beeswarm_W{W}.png'), Inches(0.0), Inches(1.3), height=Inches(5.7))
    s.shapes.add_picture(str(R / f'unified_shap/tabemb3_sofa_beeswarm_W{W}.png'), Inches(5.6), Inches(1.3), height=Inches(5.7))
    tbox(s, 0.0, 7.0, 5.4, 0.3, f'Late fusion ({W}h)', 12, 'ink', align=PP_ALIGN.CENTER); tbox(s, 5.6, 7.0, 7.6, 0.3, f'TabPFN in UTDE, 3-stream ({W}h)', 12, 'ink', align=PP_ALIGN.CENTER)

# drop every template slide except 5 (TabPFN) and 6 (UTDE); then order: setup, TabPFN, UTDE, new slides
items = list(sld)
for i in sorted([k for k in range(N_TPL) if k not in (4, 5)], reverse=True):
    prs.part.drop_rel(items[i].rId); sld.remove(items[i])
ids = list(sld); setup = ids[2]; sld.remove(setup); sld.insert(0, setup)
prs.save(OUT); print('saved', OUT, len(prs.slides), 'slides')
